from pathlib import Path
import pandas as pd
import openpyxl

REPORT = Path("reports/directory_reconciliation_report.csv")
WORKBOOK = Path("uploads/WR_DB_Ready_Final_Verified.xlsx")

if not REPORT.exists():
    raise FileNotFoundError(REPORT)

if not WORKBOOK.exists():
    raise FileNotFoundError(WORKBOOK)

# ---------------------------------------------------------
# Read all stations from workbook
# ---------------------------------------------------------

wb = openpyxl.load_workbook(WORKBOOK, read_only=True)

all_stations = set()

# Parent organizations
org_ws = wb["organizations"]

for row in org_ws.iter_rows(min_row=2, values_only=True):
    _, name, *_ = row
    if name:
        all_stations.add(str(name).strip())

# Sub-organizations
sub_ws = wb["sub_organizations"]

for row in sub_ws.iter_rows(min_row=2, values_only=True):
    _, _, name, *_ = row
    if name:
        all_stations.add(str(name).strip())

# ---------------------------------------------------------
# Read reconciliation report
# ---------------------------------------------------------

df = pd.read_csv(REPORT)

issue_stations = (
    df["Station"]
    .dropna()
    .astype(str)
    .str.strip()
)

issue_stations = set(issue_stations)

# ---------------------------------------------------------
# Perfect stations
# ---------------------------------------------------------

perfect = sorted(all_stations - issue_stations)
issues = sorted(issue_stations)

print("=" * 70)
print("Station Summary")
print("=" * 70)
print(f"Total stations          : {len(all_stations)}")
print(f"Stations with issues    : {len(issue_stations)}")
print(f"Perfect stations        : {len(perfect)}")
print("=" * 70)

print("\nPerfect Stations\n")

for i, station in enumerate(perfect, 1):
    print(f"{i:3}. {station}")

# ---------------------------------------------------------
# Save CSVs
# ---------------------------------------------------------

pd.DataFrame({"Station": perfect}).to_csv(
    "reports/perfect_stations.csv",
    index=False,
)

pd.DataFrame({"Station": issues}).to_csv(
    "reports/stations_with_issues.csv",
    index=False,
)

print("\nCreated:")
print(" - reports/perfect_stations.csv")
print(" - reports/stations_with_issues.csv")