from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl
from docx import Document


ROOT = Path(r"C:\Users\so.a\Telephone_Directory")
UPLOADS = ROOT / "uploads"
INPUT_XLSX = UPLOADS / "WR_DB_Ready_Final (1).xlsx"
INPUT_DOCX = UPLOADS / "Western Region Phone Directory 2025 Main_Telephone.docx"
OUTPUT_DIR = ROOT / "outputs" / "directory_correction"
OUTPUT_JSON = OUTPUT_DIR / "corrections.json"


def clean(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\xa0", " ").replace("\u200b", " ")
    return re.sub(r"\s+", " ", text).strip()


def norm_text(value: object) -> str:
    text = clean(value).lower().replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def norm_name(value: object) -> str:
    text = clean(value).lower()
    text = re.sub(r"\b(shri|smt|ms|mrs|mr|dr|prof|kumari|er)\.?\b", " ", text)
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def is_header(cells: list[str]) -> bool:
    text = " | ".join(cells).lower()
    return "name" in text and ("designation" in text or "mail" in text or "email" in text)


def header_key(value: str) -> str:
    text = clean(value).lower()
    if "name" in text:
        return "name"
    if "designation" in text or text in {"desg.", "desg"}:
        return "designation"
    if "department" in text:
        return "department"
    if "mobile" in text or "cell" in text:
        return "mobile"
    if "e-mail" in text or "email" in text or "mail id" in text or "e mail" in text:
        return "email"
    if "fax" in text:
        return "fax"
    if "telephone (r" in text or "residence" in text:
        return "residence"
    if "telephone" in text or "office" in text or "phone" in text:
        return "office"
    if "sl" in text or "sr" in text or "s no" in text or "s.no" in text:
        return "serial"
    return text[:25]


def looks_like_section(name: str, designation: str) -> bool:
    n = clean(name)
    d = clean(designation)
    if not n:
        return True
    if n.lower().startswith("name "):
        return True
    if n.isupper() and len(n.split()) > 2 and norm_text(n) == norm_text(d):
        return True
    section_terms = [
        "limited",
        "corporation",
        "company",
        "centre",
        "department",
        "substation",
        "switchyard",
    ]
    return norm_text(n) == norm_text(d) and any(term in n.lower() for term in section_terms)


def extract_word_records() -> list[dict]:
    doc = Document(INPUT_DOCX)
    records: list[dict] = []

    for table_index, table in enumerate(doc.tables):
        # Table 0 is the synopsis and repeats selected main-directory rows.
        if table_index == 0:
            continue

        rows = [[clean(cell.text) for cell in row.cells] for row in table.rows]
        header_index = None
        for idx, row in enumerate(rows[:15]):
            if is_header(row):
                header_index = idx
                break
        if header_index is None:
            continue

        headers: list[str] = []
        seen: dict[str, int] = defaultdict(int)
        for value in rows[header_index]:
            key = header_key(value)
            seen[key] += 1
            if seen[key] > 1 and key in {
                "name",
                "designation",
                "office",
                "residence",
                "mobile",
                "email",
                "fax",
            }:
                key = f"{key}_{seen[key]}"
            headers.append(key)

        context_parts: list[str] = []
        for row in rows[:header_index]:
            unique = []
            for value in row:
                if value and value not in unique:
                    unique.append(value)
            if unique:
                context_parts.append(" | ".join(unique))
        context = " / ".join(context_parts[-3:])

        for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            mapped: dict[str, str] = {}
            for key, value in zip(headers, row):
                if not value:
                    continue
                if key in mapped and mapped[key] != value:
                    mapped[key] = f"{mapped[key]} / {value}"
                else:
                    mapped[key] = value

            name = mapped.get("name") or mapped.get("name_2") or ""
            designation = mapped.get("designation") or mapped.get("designation_2") or ""
            if looks_like_section(name, designation) or len(name) > 100:
                continue

            record = {
                "source_table": table_index,
                "source_row": row_index,
                "context": context,
                "name": clean(name),
                "designation": clean(designation),
            }
            for key, value in mapped.items():
                if key not in {"serial", "name", "name_2", "designation", "designation_2"}:
                    record[key] = clean(value)
            records.append(record)

    return records


EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)


def extract_emails(record: dict) -> list[str]:
    found: list[str] = []
    for key, value in record.items():
        if key.startswith("email"):
            for email in EMAIL_RE.findall(value):
                email = email.strip(" .,;")
                if email and email not in found:
                    found.append(email)
    return found


def split_phone_text(value: str, default_type: str) -> list[tuple[str, str]]:
    value = clean(value)
    if not value or norm_text(value) in {"na", "n a", "nil"} or value in {"-", "--"}:
        return []

    value = EMAIL_RE.sub(" ", value)
    value = re.sub(r"\b(Extn?|Extension)\.?\s*", "Extn. ", value, flags=re.I)
    parts: list[tuple[str, str]] = []

    fax_split = re.split(r"\bFax(?:\s*No\.?)?\s*:?", value, flags=re.I)
    if len(fax_split) > 1:
        parts.extend(split_phone_text(fax_split[0], default_type))
        for fax_part in fax_split[1:]:
            parts.extend(split_phone_text(fax_part, "fax"))
        return parts

    for token in re.split(r"\s*(?:,|;|\bor\b|/)\s*", value):
        token = clean(token)
        token = token.strip(" -.;,")
        if not token or token in {"-", "--"}:
            continue
        if not re.search(r"\d", token):
            continue
        if token not in {phone for _, phone in parts}:
            parts.append((default_type, token))
    return parts


def extract_phones(record: dict) -> list[tuple[str, str]]:
    phones: list[tuple[str, str]] = []
    for key, value in record.items():
        if key.startswith("mobile"):
            candidates = split_phone_text(value, "mobile")
        elif key.startswith("fax"):
            candidates = split_phone_text(value, "fax")
        elif key.startswith("residence"):
            candidates = split_phone_text(value, "residence")
        elif key.startswith("office"):
            candidates = split_phone_text(value, "office")
        else:
            continue
        for item in candidates:
            if item not in phones:
                phones.append(item)
    return phones


def choose_record(
    employee: tuple,
    candidates: list[dict],
    occurrence_index: int,
    occurrence_count: int,
) -> tuple[dict | None, str]:
    _, _, _, name, designation = employee
    if not candidates:
        return None, "missing"
    if len(candidates) == 1:
        return candidates[0], "unique-name"

    if len(candidates) == occurrence_count:
        return candidates[occurrence_index], "name-occurrence"

    same_designation = [r for r in candidates if norm_text(r.get("designation")) == norm_text(designation)]
    if len(same_designation) == occurrence_count:
        return same_designation[occurrence_index], "designation-occurrence"
    if len(same_designation) == 1:
        return same_designation[0], "name-and-designation"

    with_email_or_phone = [
        r for r in candidates if extract_emails(r) or extract_phones(r)
    ]
    if len(with_email_or_phone) == 1:
        return with_email_or_phone[0], "only-contact-row"

    if occurrence_index < len(candidates):
        return candidates[occurrence_index], "partial-occurrence"

    return candidates[0], "ambiguous-first"


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    word_records = extract_word_records()
    records_by_name: dict[str, list[dict]] = defaultdict(list)
    for record in word_records:
        records_by_name[norm_name(record["name"])].append(record)

    wb = openpyxl.load_workbook(INPUT_XLSX, read_only=True, data_only=True)

    employees_header = [cell.value for cell in next(wb["employees"].iter_rows(min_row=1, max_row=1))]
    employees = [
        row
        for row in wb["employees"].iter_rows(min_row=2, values_only=True)
        if row and row[0] is not None
    ]
    employee_name_counts: dict[str, int] = defaultdict(int)
    for employee in employees:
        employee_name_counts[norm_name(employee[3])] += 1
    employee_seen_counts: dict[str, int] = defaultdict(int)

    corrected_employees: list[list] = []
    email_rows: list[list] = []
    phone_rows: list[list] = []
    audit_rows = [
        [
            "employee_id",
            "match_status",
            "source_table",
            "source_row",
            "old_name",
            "new_name",
            "old_designation",
            "new_designation",
            "email_count",
            "phone_count",
        ]
    ]

    email_id = 1
    phone_id = 1
    status_counts: dict[str, int] = defaultdict(int)

    for employee in employees:
        employee_id, org_id, suborg_id, old_name, old_designation = employee
        normalized_name = norm_name(old_name)
        candidates = records_by_name.get(normalized_name, [])
        occurrence_index = employee_seen_counts[normalized_name]
        employee_seen_counts[normalized_name] += 1
        record, status = choose_record(
            employee,
            candidates,
            occurrence_index,
            employee_name_counts[normalized_name],
        )
        status_counts[status] += 1

        new_name = clean(record["name"]) if record else clean(old_name)
        new_designation = clean(record["designation"]) if record else clean(old_designation)
        corrected_employees.append([employee_id, org_id, suborg_id, new_name, new_designation])

        emails = extract_emails(record) if record else []
        phones = extract_phones(record) if record else []
        for email in emails:
            email_rows.append([email_id, employee_id, email])
            email_id += 1
        for phone_type, phone_number in phones:
            phone_rows.append([phone_id, employee_id, phone_type, phone_number])
            phone_id += 1

        audit_rows.append(
            [
                employee_id,
                status,
                record.get("source_table") if record else None,
                record.get("source_row") if record else None,
                clean(old_name),
                new_name,
                clean(old_designation),
                new_designation,
                len(emails),
                len(phones),
            ]
        )

    payload = {
        "employeesHeader": employees_header,
        "employees": corrected_employees,
        "emailHeader": ["email_id  [PK]", "employee_id  [FK]", "email"],
        "emails": email_rows,
        "phoneHeader": ["phone_id  [PK]", "employee_id  [FK]", "phone_type", "phone_number"],
        "phones": phone_rows,
        "audit": audit_rows,
        "summary": {
            "wordRecords": len(word_records),
            "employees": len(corrected_employees),
            "emails": len(email_rows),
            "phones": len(phone_rows),
            "matchStatusCounts": dict(status_counts),
        },
    }

    OUTPUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload["summary"], indent=2))


if __name__ == "__main__":
    main()
