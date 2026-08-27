"""Excel generation for the Directory Versions module. Pure function: given
a DirectoryVersion and a snapshot dict from
services.directory_version_service.build_directory_snapshot(), returns raw
.xlsx bytes. Reuses the same header styling convention as every other
in-app Excel export (bold white text on navy fill).
"""

import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from utils.logger import get_logger

logger = get_logger(__name__)

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1a3a6b")


def _write_header(ws, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws, widths=None, width=22):
    # wrap_text=True on every data cell -- without it, a long value (a
    # duplicated-phone string, a multi-address email) with no natural
    # word-wrap either clips at the cell boundary (if its neighbor has
    # content) or visually spills into a blank neighbor. Not the same
    # destructive "text drawn on top of text" failure as the docx/PDF bug
    # (Excel never overlaps two cells' own text), but still a real
    # readability gap worth fixing while checking for the same bug class.
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[col[0].column_letter].width = (widths or {}).get(i, width)


def _employees_sheet(wb, title, employees):
    ws = wb.create_sheet(title[:31])
    _write_header(ws, ["Name", "Designation", "Organization", "Office Phone", "Mobile", "Email"])
    for row, e in enumerate(employees, 2):
        ws.cell(row=row, column=1, value=e.employee_name)
        ws.cell(row=row, column=2, value=e.designation or "")
        ws.cell(row=row, column=3, value=e.organization.organization_name if e.organization else "")
        ws.cell(row=row, column=4, value=e.office_phone or "")
        ws.cell(row=row, column=5, value=e.mobile_phone or "")
        ws.cell(row=row, column=6, value=e.email or "")
    _autosize(ws, widths={1: 22, 2: 22, 3: 28, 4: 20, 5: 16, 6: 28})


def _heads_sheet(wb, title, heads):
    ws = wb.create_sheet(title[:31])
    _write_header(ws, ["Name", "Designation", "Organization", "Mobile", "Email"])
    for row, h in enumerate(heads, 2):
        ws.cell(row=row, column=1, value=h.resolved_name or "")
        ws.cell(row=row, column=2, value=getattr(h, "role_title", None) or getattr(h, "designation", None) or "")
        ws.cell(row=row, column=3, value=h.organization.organization_name if h.organization else "")
        ws.cell(row=row, column=4, value=h.resolved_mobile_phone or "")
        ws.cell(row=row, column=5, value=h.resolved_email or "")
    _autosize(ws, widths={1: 22, 2: 22, 3: 28, 4: 16, 5: 28})


def _numbers_sheet(wb, title, numbers):
    ws = wb.create_sheet(title[:31])
    _write_header(ws, ["Organization", "Name", "Phone", "Email"])
    for row, n in enumerate(numbers, 2):
        ws.cell(row=row, column=1, value=n.organization or "")
        ws.cell(row=row, column=2, value=n.name or "")
        ws.cell(row=row, column=3, value=n.phone_number or "")
        ws.cell(row=row, column=4, value=n.email or "")
    _autosize(ws, widths={1: 28, 2: 22, 3: 24, 4: 28})


def _emergency_sheet(wb, contacts):
    ws = wb.create_sheet("Emergency Contacts")
    _write_header(ws, ["Employee", "Contact Name", "Relation", "Phone"])
    for row, c in enumerate(contacts, 2):
        ws.cell(row=row, column=1, value=c.employee.employee_name if c.employee else "")
        ws.cell(row=row, column=2, value=c.contact_name or "")
        ws.cell(row=row, column=3, value=c.relation or "")
        ws.cell(row=row, column=4, value=c.phone or "")
    _autosize(ws, widths={1: 22, 2: 22, 3: 16, 4: 18})


def generate_excel(version, snapshot):
    """Returns raw .xlsx bytes for `version` using `snapshot` (see
    build_directory_snapshot). Sheet order mirrors the PDF's section order."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # replaced by named sheets below

    all_employees = [e for org in snapshot["grouped"].values() for e in org["employees"]]
    all_control_rooms = [n for org in snapshot["grouped"].values() for n in org["control_rooms"]]
    all_switchyards = [n for org in snapshot["grouped"].values() for n in org["switchyards"]]

    _employees_sheet(wb, "Employees", all_employees)
    _employees_sheet(wb, "Utility Heads", snapshot["utility_heads"])
    _heads_sheet(wb, "Administrative Heads", snapshot["administrative_heads"])
    _employees_sheet(wb, "KMP", snapshot["kmp"])
    _numbers_sheet(wb, "Control Rooms", all_control_rooms)
    _numbers_sheet(wb, "Switchyards", all_switchyards)
    _emergency_sheet(wb, snapshot["emergency_contacts"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
