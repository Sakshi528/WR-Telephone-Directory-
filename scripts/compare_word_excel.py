"""Compares Word directory rosters against the verified Excel workbook --
employees (WRLDC + region-wide KMP), control rooms, and switchyards, per
station. Read-only throughout. Writes reports/word_vs_excel_2026.csv in the
same Station/Category/Issue/Name/Details schema as
reports/directory_reconciliation_report.csv (scripts/reconcile_word_database.py),
so scripts/filter_reconciliation_artifacts.py's pairing logic applies to
both unchanged.

Segmentation, station matching, and record grouping mirror
reconcile_word_database.py as closely as the two data sources allow -- see
that script's own docstring for the shared Word-parsing approach. Control
rooms and switchyards are compared by actual label (cr_label in the
workbook), not just a per-station present/absent flag, matching how the
database-side comparison already works, since without a per-record name
there is nothing for the naming-artifact filter to pair against.

Usage:
  python scripts/compare_word_excel.py
"""

import csv
import os
import re
import sys
from collections import Counter, defaultdict

import docx
import openpyxl
from docx.table import Table
from docx.text.paragraph import Paragraph

WORD_PATH = "uploads/Western Region Phone Directory 2026 Main_TD.docx"
EXCEL_PATH = "uploads/WR_DB_Ready_Final_Verified_v3.xlsx"
OUTPUT_CSV = "reports/word_vs_excel_2026.csv"

NUMBERED_TITLE_RE = re.compile(r"^\d+\.\s*[A-Za-z]")
LEADING_NUMBER_RE = re.compile(r"^\d+(\.\d+)*\.?\s*")

# matches both "switchyard" and "switch yard" (Word uses both spellings)
SWITCHYARD_RE = re.compile(r"switch\s*yard", re.IGNORECASE)

# content-based (not column-position) classifiers for phone/email cell values
PHONE_RE = re.compile(r"\d{5,}")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.\w+")


def clean(v):
    return re.sub(r"\s+", " ", str(v or "").replace("\xa0", " ")).strip()


def strip_leading_number(text):
    return LEADING_NUMBER_RE.sub("", text).strip()


def dedup_preserve_order(seq):
    seen = set()
    out = []
    for v in seq:
        if v not in seen:
            seen.add(v)
            out.append(v)
    return out


def classify_row_content(cells):
    phones, emails, others = [], [], []
    for raw in cells:
        v = clean(raw)
        if not v:
            continue
        if EMAIL_RE.search(v):
            emails.append(v)
        elif PHONE_RE.search(v):
            phones.append(v)
        else:
            others.append(v)
    return dedup_preserve_order(phones), dedup_preserve_order(emails), dedup_preserve_order(others)


def normalize_phone_set(text):
    return frozenset(PHONE_RE.findall(text or ""))


def normalize_email_set(text):
    return frozenset(e.lower() for e in EMAIL_RE.findall(text or ""))


# ─── EXCEL LOADING ───────────────────────────────────────────────────────────

def load_excel(path):
    """v3 workbook sheets, per the actual current headers (verified against
    the live file, not assumed from the older raw workbook):
      organizations       : org_id, name, address, state
      sub_organizations   : suborg_id, org_id, name, address, state
      employees           : employee_id, org_id, suborg_id, department, name, designation
      KMP                 : kmp_id, organization_id, name, designation, office_phone,
                             residence_phone, mobile_phone, email  -- org-level only,
                             no suborg linkage, so KMP contacts are bucketed under the
                             parent organization, never a specific sub-organization.
      control_rooms       : id, org_id, suborg_id, organization_name, sub_organization_name,
                             cr_label, designation, phone_1, phone_2, mobile_1, mobile_2, email_1
      switchyards          : id, org_id, suborg_id, organization_name, sub_organization_name,
                             cr_label, phone_1, mobile_1, email_1
    """
    wb = openpyxl.load_workbook(path, read_only=True)

    def data_rows(name):
        return list(wb[name].iter_rows(values_only=True))[1:]

    org_name_by_id = {}
    org_id_by_name_lower = {}
    for oid, name, _addr, _state in data_rows("organizations"):
        name = clean(name)
        if oid is None or not name:
            continue
        org_name_by_id[oid] = name
        org_id_by_name_lower[name.lower()] = oid

    suborg_name_by_id = {}
    suborg_id_by_name_lower = {}
    for sid, _parent_oid, name, _addr, _state in data_rows("sub_organizations"):
        name = strip_leading_number(clean(name))
        if sid is None or not name:
            continue
        suborg_name_by_id[sid] = name
        suborg_id_by_name_lower[name.lower()] = sid

    emp_by_suborg = defaultdict(list)
    emp_by_org_direct = defaultdict(list)
    for emp_id, org_id, suborg_id, _department, name, _designation in data_rows("employees"):
        name = clean(name)
        if not name:
            continue
        if suborg_id is not None:
            emp_by_suborg[suborg_id].append(name)
        elif org_id is not None:
            emp_by_org_direct[org_id].append(name)

    # KMP has no suborg_id column -- every KMP contact is bucketed at the
    # parent organization level only, merged into the same org-level index
    # employees already populate (a Word station matched at "org" level sees
    # both its own WRLDC-style staff and its KMP leadership together).
    for _kmp_id, org_id, name, _designation, *_contact in data_rows("KMP"):
        name = clean(name)
        if not name or org_id is None:
            continue
        emp_by_org_direct[org_id].append(name)

    cr_by_suborg = defaultdict(list)
    cr_by_org_direct = defaultdict(list)
    for _id, org_id, suborg_id, _org_name, _suborg_name, cr_label, *_rest in data_rows("control_rooms"):
        label = clean(cr_label)
        if not label:
            continue
        if suborg_id is not None:
            cr_by_suborg[suborg_id].append(label)
        elif org_id is not None:
            cr_by_org_direct[org_id].append(label)

    sw_by_suborg = defaultdict(list)
    sw_by_org_direct = defaultdict(list)
    for _id, org_id, suborg_id, _org_name, _suborg_name, cr_label, *_rest in data_rows("switchyards"):
        label = clean(cr_label)
        if not label:
            continue
        if suborg_id is not None:
            sw_by_suborg[suborg_id].append(label)
        elif org_id is not None:
            sw_by_org_direct[org_id].append(label)

    return {
        "org_name_by_id": org_name_by_id,
        "org_id_by_name_lower": org_id_by_name_lower,
        "suborg_name_by_id": suborg_name_by_id,
        "suborg_id_by_name_lower": suborg_id_by_name_lower,
        "emp_by_suborg": emp_by_suborg,
        "emp_by_org_direct": emp_by_org_direct,
        "cr_by_suborg": cr_by_suborg,
        "cr_by_org_direct": cr_by_org_direct,
        "sw_by_suborg": sw_by_suborg,
        "sw_by_org_direct": sw_by_org_direct,
    }


# ─── WORD PARSING ────────────────────────────────────────────────────────────

def iter_block_items(doc):
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def find_header_columns(rows_cells):
    for cells in rows_cells:
        name_col = next((i for i, c in enumerate(cells) if "name" in c.lower()), None)
        desig_col = next((i for i, c in enumerate(cells) if "designat" in c.lower()), None)
        if name_col is not None and desig_col is not None:
            return name_col, desig_col
    return None, None


def is_header_row(cells):
    lower = [c.lower() for c in cells]
    return any("name" in c for c in lower) and any("designat" in c for c in lower)


def pick_label(name_cell, desig_cell, bare_words):
    """Prefers whichever cell is NOT just the bare category word (e.g.
    "SLDC DNH" / "Control room" -> "SLDC DNH"), so distinctly-named rows
    aren't collapsed into a shared, useless label that then falsely
    mismatches the workbook's specific cr_label -- same rule
    reconcile_word_database.py uses for the database-side comparison."""
    name_bare = name_cell.strip().lower() in bare_words
    desig_bare = desig_cell.strip().lower() in bare_words
    if desig_bare and not name_bare and name_cell.strip():
        return name_cell
    if name_bare and not desig_bare and desig_cell.strip():
        return desig_cell
    return name_cell if name_cell.strip() else desig_cell


def parse_table_segments(table, fallback_name):
    """Split one Word table into (name, employee_set, cr_labels, sw_labels)
    segments. Switchyard rows are checked *before* the generic "control
    room" substring check (via SWITCHYARD_RE) -- the previous version of
    this function only checked for "control room" and silently mis-filed
    every switchyard row as an employee name."""
    rows_cells = [[clean(c.text) for c in r.cells] for r in table.rows]
    name_col, desig_col = find_header_columns(rows_cells)
    if name_col is None:
        return []  # not a roster table (reference/appendix table) -- skip

    segments = []
    current_name = None
    current_employees = {}
    current_cr = []
    current_sw = []

    def flush():
        if current_employees or current_cr or current_sw:
            segments.append((
                current_name or fallback_name,
                set(current_employees.values()),
                list(current_cr),
                list(current_sw),
            ))

    for cells in rows_cells:
        if not cells:
            continue
        first_cell = cells[0]

        if NUMBERED_TITLE_RE.match(first_cell):
            flush()
            current_name = strip_leading_number(first_cell)
            current_employees = {}
            current_cr = []
            current_sw = []
            continue

        if is_header_row(cells):
            continue  # roster header, or a repeated print-pagination header

        if not any(cells):
            continue

        name_cell = cells[name_col] if name_col < len(cells) else ""
        desig_cell = cells[desig_col] if desig_col < len(cells) else ""

        if not name_cell:
            continue

        if SWITCHYARD_RE.search(name_cell) or SWITCHYARD_RE.search(desig_cell):
            current_sw.append(pick_label(name_cell, desig_cell, {"switchyard", "switch yard"}))
            continue

        if "control room" in name_cell.lower() or "control room" in desig_cell.lower():
            current_cr.append(pick_label(name_cell, desig_cell, {"control room", "control rooms"}))
            continue

        if name_cell.lower() == desig_cell.lower():
            continue  # e.g. "GUJARAT" / "GUJARAT" title-block noise row

        current_employees[name_cell.lower()] = name_cell

    flush()
    return segments


def parse_word_segments(path):
    doc = docx.Document(path)
    current_heading = None
    all_segments = []
    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            if item.style.name.startswith("Heading") and clean(item.text):
                current_heading = strip_leading_number(clean(item.text))
            continue
        all_segments.extend(parse_table_segments(item, current_heading))
    return all_segments


# ─── MATCHING WORD SEGMENTS TO EXCEL ORGANIZATIONS ──────────────────────────

def _best_substring_match(key_lower, name_lower_to_id):
    """Fallback for names like 'Kawas' that should match 'NTPC Kawas' --
    picks the closest-length known name containing (or contained by) it."""
    if len(key_lower) < 4:
        return None
    best = None
    best_diff = None
    for known_lower, ident in name_lower_to_id.items():
        if known_lower in key_lower or key_lower in known_lower:
            diff = abs(len(known_lower) - len(key_lower))
            if best_diff is None or diff < best_diff:
                best, best_diff = ident, diff
    return best


def match_segment(name, excel):
    key_lower = strip_leading_number(name).lower()

    if key_lower in excel["suborg_id_by_name_lower"]:
        sid = excel["suborg_id_by_name_lower"][key_lower]
        return ("suborg", sid, excel["suborg_name_by_id"][sid])
    if key_lower in excel["org_id_by_name_lower"]:
        oid = excel["org_id_by_name_lower"][key_lower]
        return ("org", oid, excel["org_name_by_id"][oid])

    sid = _best_substring_match(key_lower, excel["suborg_id_by_name_lower"])
    if sid is not None:
        return ("suborg", sid, excel["suborg_name_by_id"][sid])
    oid = _best_substring_match(key_lower, excel["org_id_by_name_lower"])
    if oid is not None:
        return ("org", oid, excel["org_name_by_id"][oid])

    return ("none", key_lower, name)


def group_segments(segments, excel):
    groups = {}
    for name, employees, cr_labels, sw_labels in segments:
        match_type, match_id, display_name = match_segment(name, excel)
        key = (match_type, match_id)
        g = groups.setdefault(key, {
            "display_name": display_name,
            "employees": {},
            "cr_labels": [],
            "sw_labels": [],
            "match_type": match_type,
            "match_id": match_id,
        })
        for emp in employees:
            g["employees"][emp.lower()] = emp
        g["cr_labels"].extend(cr_labels)
        g["sw_labels"].extend(sw_labels)
    return groups


# ─── RECORD COMPARISON (Missing / Extra / Duplicate, by label) ─────────────

def compare_labels(station, category, word_labels, excel_labels, totals, cat_key):
    """word_labels/excel_labels: lists of raw label strings for this
    station+category. Missing/Extra determined by exact case-insensitive
    match; a label appearing more than once on the Excel side is flagged
    Duplicate. No phone/email comparison here -- the workbook's
    control_rooms/switchyards sheets aren't compared field-by-field in this
    pass, only by whether the record itself is present on both sides."""
    rows = []

    word_by_lower = {}
    for label in word_labels:
        word_by_lower.setdefault(label.strip().lower(), label)

    excel_counts = Counter(label.strip().lower() for label in excel_labels)
    excel_by_lower = {}
    for label in excel_labels:
        excel_by_lower.setdefault(label.strip().lower(), label)

    word_keys = set(word_by_lower)
    excel_keys = set(excel_by_lower)

    for key in sorted(word_keys - excel_keys):
        rows.append({
            "Station": station, "Category": category, "Issue": "Missing",
            "Name": word_by_lower[key],
            "Details": f"Word lists a {category.lower()} here; none found in Excel workbook",
        })
        totals[f"missing_{cat_key}"] += 1

    for key in sorted(excel_keys - word_keys):
        rows.append({
            "Station": station, "Category": category, "Issue": "Extra",
            "Name": excel_by_lower[key],
            "Details": f"Excel workbook has a {category.lower()} here; none found in Word directory",
        })
        totals[f"extra_{cat_key}"] += 1

    for key, count in excel_counts.items():
        if count > 1:
            rows.append({
                "Station": station, "Category": category, "Issue": "Duplicate",
                "Name": excel_by_lower[key],
                "Details": f"{count} {category.lower()} records share this label in the Excel workbook",
            })
            totals[f"duplicate_{cat_key}"] += 1

    return rows


def build_report_rows(groups, excel):
    rows = []
    totals = defaultdict(int)
    skipped_non_directory = []

    for (match_type, match_id), g in groups.items():
        if match_type == "none":
            skipped_non_directory.append(g["display_name"])
            continue

        station = g["display_name"]
        word_names = g["employees"]  # lower -> original

        if match_type == "suborg":
            excel_names = excel["emp_by_suborg"].get(match_id, [])
            excel_cr = excel["cr_by_suborg"].get(match_id, [])
            excel_sw = excel["sw_by_suborg"].get(match_id, [])
        else:  # "org"
            excel_names = excel["emp_by_org_direct"].get(match_id, [])
            excel_cr = excel["cr_by_org_direct"].get(match_id, [])
            excel_sw = excel["sw_by_org_direct"].get(match_id, [])

        excel_names_lower = {}
        for n in excel_names:
            excel_names_lower.setdefault(n.strip().lower(), n)

        missing = sorted(orig for lower, orig in word_names.items() if lower not in excel_names_lower)
        extra = sorted(excel_names_lower[lower] for lower in excel_names_lower if lower not in word_names)

        for name in missing:
            rows.append({
                "Station": station, "Category": "Employee", "Issue": "Missing",
                "Name": name, "Details": "In Word directory but not found in Excel workbook",
            })
            totals["missing_emp"] += 1
        for name in extra:
            rows.append({
                "Station": station, "Category": "Employee", "Issue": "Extra",
                "Name": name, "Details": "In Excel workbook but not found in Word directory",
            })
            totals["extra_emp"] += 1

        rows.extend(compare_labels(station, "Control Room", g["cr_labels"], excel_cr, totals, "cr"))
        rows.extend(compare_labels(station, "Switchyard", g["sw_labels"], excel_sw, totals, "sw"))

    rows.sort(key=lambda r: (r["Station"].lower(), r["Category"], r["Issue"]))
    return rows, totals, skipped_non_directory


def run(word_path, excel_path, output_csv):
    excel = load_excel(excel_path)
    segments = parse_word_segments(word_path)
    groups = group_segments(segments, excel)
    rows, totals, skipped = build_report_rows(groups, excel)

    os.makedirs(os.path.dirname(output_csv), exist_ok=True)
    with open(output_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["Station", "Category", "Issue", "Name", "Details"])
        writer.writeheader()
        writer.writerows(rows)

    print("=" * 70)
    print("  Word <-> Excel (v3) Comparison Summary")
    print("=" * 70)
    print(f"  Non-directory / unmatched sections ignored : {len(skipped)}")
    for name in skipped:
        print(f"    - {name}")
    print(f"  Total rows written     : {len(rows)}")
    print(f"  Missing employees      : {totals['missing_emp']}")
    print(f"  Extra employees        : {totals['extra_emp']}")
    print(f"  Missing control rooms  : {totals['missing_cr']}")
    print(f"  Extra control rooms    : {totals['extra_cr']}")
    print(f"  Duplicate control rooms: {totals['duplicate_cr']}")
    print(f"  Missing switchyards    : {totals['missing_sw']}")
    print(f"  Extra switchyards      : {totals['extra_sw']}")
    print(f"  Duplicate switchyards  : {totals['duplicate_sw']}")
    print("=" * 70)
    print(f"  Report written to: {output_csv}")


if __name__ == "__main__":
    run(WORD_PATH, EXCEL_PATH, OUTPUT_CSV)
