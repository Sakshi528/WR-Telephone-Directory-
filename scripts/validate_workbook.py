"""
Validates the source workbook before it is fed to import_from_excel_db.py.

Read-only: does not touch the database and does not import/modify any
other script. Checks the organizations, sub_organizations, employees,
control_rooms, and switchyards sheets for:

  - duplicate primary-key IDs
  - blank required fields
  - orphan / invalid org_id and suborg_id references

Usage:
  python scripts/validate_workbook.py
  python scripts/validate_workbook.py path/to/other_workbook.xlsx
"""

import sys
from collections import Counter

import openpyxl

EXCEL_PATH = "uploads/WR_DB_Ready_Final.xlsx"

REQUIRED_SHEETS = [
    "organizations",
    "sub_organizations",
    "employees",
    "control_rooms",
    "switchyards",
]


class Issue:
    def __init__(self, severity, sheet, message):
        self.severity = severity  # "ERROR" | "WARNING"
        self.sheet = sheet
        self.message = message


def clean(v):
    return str(v).strip() if v is not None else ""


def sheet_data_rows(wb, name):
    return list(wb[name].iter_rows(values_only=True))[1:]


def validate_organizations(rows, issues):
    seen = Counter()
    org_ids = set()
    for i, r in enumerate(rows, start=2):
        oid, name = r[0], r[1]
        if oid is None:
            issues.append(Issue("ERROR", "organizations", f"Row {i}: missing org_id"))
            continue
        seen[oid] += 1
        org_ids.add(oid)
        if not clean(name):
            issues.append(Issue("ERROR", "organizations", f"Row {i}: org_id={oid} has blank name"))
    for oid, count in seen.items():
        if count > 1:
            issues.append(Issue("ERROR", "organizations", f"Duplicate org_id={oid} appears {count} times"))
    return org_ids


def validate_sub_organizations(rows, org_ids, issues):
    seen = Counter()
    suborg_ids = set()
    for i, r in enumerate(rows, start=2):
        sid, parent_oid, name = r[0], r[1], r[2]
        if sid is None:
            issues.append(Issue("ERROR", "sub_organizations", f"Row {i}: missing suborg_id"))
            continue
        seen[sid] += 1
        suborg_ids.add(sid)
        if not clean(name):
            issues.append(Issue("ERROR", "sub_organizations", f"Row {i}: suborg_id={sid} has blank name"))
        if parent_oid is None:
            issues.append(Issue("WARNING", "sub_organizations", f"Row {i}: suborg_id={sid} has no org_id"))
        elif parent_oid not in org_ids:
            issues.append(Issue("ERROR", "sub_organizations",
                                 f"Row {i}: suborg_id={sid} references unknown org_id={parent_oid}"))
    for sid, count in seen.items():
        if count > 1:
            issues.append(Issue("ERROR", "sub_organizations", f"Duplicate suborg_id={sid} appears {count} times"))
    return suborg_ids


def validate_employees(rows, org_ids, suborg_ids, issues):
    # Current header: employee_id [PK], org_id [FK], suborg_id [FK], name, designation
    seen = Counter()
    for i, r in enumerate(rows, start=2):
        emp_id, org_id, suborg_id, name, _designation = r[:5]
        if emp_id is None:
            issues.append(Issue("ERROR", "employees", f"Row {i}: missing employee_id"))
            continue
        seen[emp_id] += 1
        if not clean(name):
            issues.append(Issue("WARNING", "employees", f"Row {i}: employee_id={emp_id} has blank name"))
        if org_id is not None and org_id not in org_ids:
            issues.append(Issue("ERROR", "employees",
                                 f"Row {i}: employee_id={emp_id} references unknown org_id={org_id}"))
        if suborg_id is not None and suborg_id not in suborg_ids:
            issues.append(Issue("ERROR", "employees",
                                 f"Row {i}: employee_id={emp_id} references unknown suborg_id={suborg_id}"))
    for emp_id, count in seen.items():
        if count > 1:
            issues.append(Issue("ERROR", "employees", f"Duplicate employee_id={emp_id} appears {count} times"))


def validate_directory_sheet(rows, sheet_name, org_ids, suborg_ids, issues):
    seen = Counter()
    for i, r in enumerate(rows, start=2):
        rid, org_id, suborg_id = r[0], r[1], r[2]
        if rid is None:
            issues.append(Issue("WARNING", sheet_name, f"Row {i}: missing id"))
        else:
            seen[rid] += 1
        if org_id is not None and org_id not in org_ids:
            issues.append(Issue("ERROR", sheet_name, f"Row {i}: id={rid} references unknown org_id={org_id}"))
        if suborg_id is not None and suborg_id not in suborg_ids:
            issues.append(Issue("ERROR", sheet_name, f"Row {i}: id={rid} references unknown suborg_id={suborg_id}"))
    for rid, count in seen.items():
        if count > 1:
            issues.append(Issue("ERROR", sheet_name, f"Duplicate id={rid} appears {count} times"))


def print_report(issues, sheet_row_counts):
    sep = "=" * 70
    print(sep)
    print("  Workbook Validation Report")
    print(sep)

    print("\n  Rows per sheet:")
    for name, count in sheet_row_counts.items():
        print(f"    {name:<20} {count}")

    errors = [i for i in issues if i.severity == "ERROR"]
    warnings = [i for i in issues if i.severity == "WARNING"]

    for severity, group in (("ERROR", errors), ("WARNING", warnings)):
        print(f"\n  {severity}S ({len(group)}):")
        if not group:
            print("    (none)")
            continue
        by_sheet = {}
        for issue in group:
            by_sheet.setdefault(issue.sheet, []).append(issue.message)
        for sheet, messages in by_sheet.items():
            print(f"\n    [{sheet}] {len(messages)} {severity.lower()}(s)")
            for msg in messages:
                print(f"      - {msg}")

    print(f"\n{sep}")
    print(f"  TOTAL: {len(errors)} error(s), {len(warnings)} warning(s)")
    print(sep)

    return len(errors) == 0


def run(path):
    wb = openpyxl.load_workbook(path, read_only=True)

    missing_sheets = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
    if missing_sheets:
        print(f"ERROR: workbook is missing required sheet(s): {missing_sheets}")
        return False

    sheet_rows = {name: sheet_data_rows(wb, name) for name in REQUIRED_SHEETS}
    sheet_row_counts = {name: len(rows) for name, rows in sheet_rows.items()}

    issues = []
    org_ids = validate_organizations(sheet_rows["organizations"], issues)
    suborg_ids = validate_sub_organizations(sheet_rows["sub_organizations"], org_ids, issues)
    validate_employees(sheet_rows["employees"], org_ids, suborg_ids, issues)
    validate_directory_sheet(sheet_rows["control_rooms"], "control_rooms", org_ids, suborg_ids, issues)
    validate_directory_sheet(sheet_rows["switchyards"], "switchyards", org_ids, suborg_ids, issues)

    return print_report(issues, sheet_row_counts)


if __name__ == "__main__":
    excel_path = sys.argv[1] if len(sys.argv) > 1 else EXCEL_PATH
    ok = run(excel_path)
    sys.exit(0 if ok else 1)
